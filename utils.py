"""
Funciones compartidas por toda la app: lectura del Excel (por tablas
nombradas, para que funcione con archivos actualizados del mismo formato),
limpieza de datos y filtros de sidebar.

Los filtros se definen UNA sola vez, en app.py (el archivo de entrada),
que actúa como "marco" común alrededor de todas las páginas. Así su
estado se mantiene siempre al navegar, sin importar la página activa.
"""

import io
import openpyxl
import pandas as pd
import streamlit as st

# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------

PROVINCIA_LABELS = {
    "Chubut": "Chubut",
    "Rio negro": "Río Negro",
    "Santa Cruz": "Santa Cruz",
}

TIPO_FIX = {
    "autos": "Autos",
    "PIck-Up": "Pick-Up",
}

COLOR_PROVINCIA = {
    "Chubut": "#1f77b4",
    "Río Negro": "#ff7f0e",
    "Santa Cruz": "#2ca02c",
}


# ---------------------------------------------------------------------------
# Lectura del Excel por tabla nombrada (robusto a filas/columnas agregadas)
# ---------------------------------------------------------------------------

def _read_named_table(wb: openpyxl.Workbook, table_name: str) -> pd.DataFrame | None:
    """Busca una Tabla de Excel por su nombre en todas las hojas y la
    devuelve como DataFrame, usando la primera fila del rango como header."""
    for ws in wb.worksheets:
        if table_name in ws.tables:
            ref = ws.tables[table_name].ref
            data = ws[ref]
            rows = [[cell.value for cell in row] for row in data]
            if not rows:
                return None
            header, body = rows[0], rows[1:]
            return pd.DataFrame(body, columns=header)
    return None


def _clean_pat(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]

    for col in ["Localidad", "Modelo", "Marca", "tipo", "provincia", "Zona", "estado"]:
        df[col] = df[col].astype(str).str.strip()

    df["tipo"] = df["tipo"].replace(TIPO_FIX)
    df["provincia_label"] = df["provincia"].map(PROVINCIA_LABELS).fillna(df["provincia"])
    df["Periodo"] = pd.to_datetime(df["Periodo"], errors="coerce")
    df["Valor"] = pd.to_numeric(df["Valor"], errors="coerce").fillna(0).astype(int)
    df = df.dropna(subset=["Periodo"])
    return df


def _clean_aux(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    for col in df.columns:
        if df[col].dtype == object:
            df[col] = df[col].astype(str).str.strip()
    return df


@st.cache_data(show_spinner=False)
def load_workbook_data(file_bytes: bytes):
    """Carga y limpia las 3 tablas de interés a partir de los bytes del
    archivo subido. Cacheado por contenido del archivo."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    df_pat = _read_named_table(wb, "PaT_Ciudad")
    df_modelos = _read_named_table(wb, "Modelos")
    df_ciudades = _read_named_table(wb, "Ciudades")

    missing = [n for n, d in [("PaT_Ciudad", df_pat), ("Modelos", df_modelos),
                               ("Ciudades", df_ciudades)] if d is None]
    if missing:
        raise ValueError(
            "No encontré la(s) tabla(s) nombrada(s): " + ", ".join(missing) +
            ". Verificá que el archivo tenga el mismo formato (mismos nombres "
            "de Tabla de Excel: PaT_Ciudad, Modelos, Ciudades)."
        )

    df_pat = _clean_pat(df_pat)
    df_modelos = _clean_aux(df_modelos)
    df_ciudades = _clean_aux(df_ciudades)

    return df_pat, df_modelos, df_ciudades


def ensure_data_loaded():
    """Muestra el file_uploader en el sidebar (si hace falta) y garantiza
    que los dataframes estén en session_state. Se llama una sola vez desde
    app.py. Devuelve (df_pat, df_modelos, df_ciudades)."""

    with st.sidebar:
        st.markdown("### 📁 Datos")
        uploaded = st.file_uploader(
            "Excel de patentamientos (mismo formato)",
            type=["xlsx"],
            key="uploader",
            help="Debe contener las tablas 'PaT_Ciudad' (hoja Tabla Pat) y "
                 "'Modelos' / 'Ciudades' (hoja auxiliares).",
        )
        if uploaded is not None:
            try:
                file_bytes = uploaded.getvalue()
                if st.session_state.get("_file_hash") != hash(file_bytes):
                    df_pat, df_modelos, df_ciudades = load_workbook_data(file_bytes)
                    st.session_state["df_pat"] = df_pat
                    st.session_state["df_modelos"] = df_modelos
                    st.session_state["df_ciudades"] = df_ciudades
                    st.session_state["_file_hash"] = hash(file_bytes)
                    st.session_state["file_name"] = uploaded.name
                st.success(f"✅ {st.session_state.get('file_name', 'Archivo')} cargado")
            except Exception as e:
                st.error(f"Error al leer el archivo: {e}")

    if "df_pat" not in st.session_state:
        st.info("👈 Subí el Excel de patentamientos en la barra lateral para empezar.")
        st.stop()

    return (
        st.session_state["df_pat"],
        st.session_state["df_modelos"],
        st.session_state["df_ciudades"],
    )


# ---------------------------------------------------------------------------
# Filtros de sidebar — se llaman UNA vez desde app.py (marco común)
# ---------------------------------------------------------------------------

def sidebar_filters(df: pd.DataFrame):
    """Aplica los filtros del sidebar. Devuelve (df_activo, df_bruto, df_neto):
    - df_bruto: solo registros 'Total' (volumen global, sin ajustar)
    - df_neto: 'Total' + '-HERZT' (volumen real, descontando los
      patentamientos de Hertz bajo el régimen provincial vigente desde 2025,
      que en su mayoría no circulan en la provincia)
    - df_activo: el que corresponde según el selector "Volumen" — es el que
      usan Rankings, Comparativas, Evolución y Detalle.
    Los tres respetan el resto de los filtros (período, provincia, zona,
    tipo, marca, localidad)."""
    st.sidebar.markdown("### 🔎 Filtros")

    periodos = sorted(df["Periodo"].unique())
    periodo_labels = [pd.Timestamp(p).strftime("%Y-%m") for p in periodos]
    label_to_ts = dict(zip(periodo_labels, periodos))

    sel_labels = st.sidebar.select_slider(
        "Período",
        options=periodo_labels,
        value=(periodo_labels[0], periodo_labels[-1]),
        key="f_periodo",
    )
    p_ini, p_fin = label_to_ts[sel_labels[0]], label_to_ts[sel_labels[1]]

    provincias = sorted(df["provincia_label"].unique())
    sel_prov = st.sidebar.multiselect("Provincia", provincias, default=provincias, key="f_prov")

    zonas = sorted(df["Zona"].unique())
    sel_zona = st.sidebar.multiselect("Zona", zonas, default=zonas, key="f_zona")

    tipos = sorted(df["tipo"].unique())
    sel_tipo = st.sidebar.multiselect("Tipo de vehículo", tipos, default=tipos, key="f_tipo")

    marcas = sorted(df["Marca"].unique())
    sel_marca = st.sidebar.multiselect(
        "Marca (vacío = todas)", marcas, default=[], key="f_marca"
    )

    localidades = sorted(df[df["provincia_label"].isin(sel_prov)]["Localidad"].unique())
    sel_loc = st.sidebar.multiselect(
        "Localidad (vacío = todas)", localidades, default=[], key="f_loc"
    )

    base = df[
        (df["Periodo"] >= p_ini) & (df["Periodo"] <= p_fin)
        & (df["provincia_label"].isin(sel_prov))
        & (df["Zona"].isin(sel_zona))
        & (df["tipo"].isin(sel_tipo))
    ]
    if sel_marca:
        base = base[base["Marca"].isin(sel_marca)]
    if sel_loc:
        base = base[base["Localidad"].isin(sel_loc)]

    df_bruto = base[base["estado"] == "Total"]
    df_neto = base[base["estado"].isin(["Total", "-HERZT"])]

    st.sidebar.markdown("### 💰 Volumen")
    modo = st.sidebar.radio(
        "¿Qué volumen usar en los análisis?",
        ["Global (bruto)", "Real (neto de Hertz)"],
        index=0,
        key="f_modo_volumen",
        help="Desde marzo/abril de 2025, Hertz patenta vehículos en la "
             "provincia por una normativa que le resulta beneficiosa, pero "
             "la mayoría no circula realmente en la provincia. 'Real' resta "
             "esos patentamientos (registrados como -HERZT) para un análisis "
             "más transparente del volumen genuino.",
    )
    df_activo = df_bruto if modo.startswith("Global") else df_neto

    st.sidebar.caption(
        f"{len(df_activo):,} registros filtrados de {len(df):,} totales".replace(",", ".")
    )

    return df_activo, df_bruto, df_neto


# ---------------------------------------------------------------------------
# Helpers de formato
# ---------------------------------------------------------------------------

def fmt_int(n) -> str:
    return f"{int(n):,}".replace(",", ".")


def month_label(ts) -> str:
    meses = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
    ts = pd.Timestamp(ts)
    return f"{meses[ts.month - 1]}-{str(ts.year)[2:]}"


# ---------------------------------------------------------------------------
# Resumen narrativo automático (por reglas, sin servicios externos)
# ---------------------------------------------------------------------------

def generar_resumen_reglas(df: pd.DataFrame, df_bruto: pd.DataFrame = None, df_neto: pd.DataFrame = None) -> str:
    """Arma un párrafo en español a partir de los datos ya filtrados,
    sin depender de ningún servicio externo. Gratis e instantáneo.
    Si se pasan df_bruto y df_neto, agrega una mención al ajuste por
    Hertz (diferencia entre volumen global y volumen real)."""
    if df.empty:
        return "No hay datos para el período y los filtros seleccionados."

    total = int(df["Valor"].sum())
    periodos = sorted(df["Periodo"].unique())
    rango_ini, rango_fin = month_label(periodos[0]), month_label(periodos[-1])
    n_meses = len(periodos)

    monthly = df.groupby("Periodo")["Valor"].sum().sort_index()

    by_prov = df.groupby("provincia_label")["Valor"].sum().sort_values(ascending=False)
    by_marca = df.groupby("Marca")["Valor"].sum().sort_values(ascending=False)
    by_tipo = df.groupby("tipo")["Valor"].sum().sort_values(ascending=False)
    by_loc = df.groupby("Localidad")["Valor"].sum().sort_values(ascending=False)

    partes = []
    partes.append(
        f"Entre **{rango_ini}** y **{rango_fin}** ({n_meses} "
        f"{'mes' if n_meses == 1 else 'meses'}) se patentaron **{fmt_int(total)}** "
        f"vehículos dentro de los filtros seleccionados."
    )

    if total > 0 and len(by_prov) and len(by_loc):
        prov_lider, prov_pct = by_prov.index[0], by_prov.iloc[0] / total * 100
        loc_lider = by_loc.index[0]
        partes.append(
            f"**{prov_lider}** concentra la mayor participación, con "
            f"{prov_pct:.0f}% del total, y **{loc_lider}** es la localidad con más "
            f"patentamientos dentro de ese recorte."
        )

    if total > 0 and len(by_marca) and len(by_tipo):
        marca_lider, marca_pct = by_marca.index[0], by_marca.iloc[0] / total * 100
        tipo_lider, tipo_pct = by_tipo.index[0], by_tipo.iloc[0] / total * 100
        partes.append(
            f"La marca más patentada fue **{marca_lider}** ({marca_pct:.0f}% del "
            f"total), y el tipo de vehículo dominante fue **{tipo_lider}** "
            f"({tipo_pct:.0f}%)."
        )

    if n_meses >= 2:
        mes_pico, mes_valle = monthly.idxmax(), monthly.idxmin()
        partes.append(
            f"El mes de mayor actividad fue **{month_label(mes_pico)}** "
            f"({fmt_int(monthly.max())} unidades), mientras que "
            f"**{month_label(mes_valle)}** registró el menor volumen "
            f"({fmt_int(monthly.min())})."
        )
        primero, ultimo = monthly.iloc[0], monthly.iloc[-1]
        if primero > 0:
            var = (ultimo - primero) / primero * 100
            direccion = "creció" if var > 0.5 else "cayó" if var < -0.5 else "se mantuvo estable"
            partes.append(
                f"Comparando el primer y el último mes del rango, el volumen "
                f"{direccion}{'' if direccion == 'se mantuvo estable' else f' un {abs(var):.0f}%'}."
            )

    if df_bruto is not None and df_neto is not None and not df_bruto.empty:
        bruto_total = int(df_bruto["Valor"].sum())
        neto_total = int(df_neto["Valor"].sum())
        ajuste = bruto_total - neto_total
        if ajuste > 0 and bruto_total > 0:
            pct_ajuste = ajuste / bruto_total * 100
            partes.append(
                f"Del volumen global (**{fmt_int(bruto_total)}**), **{fmt_int(ajuste)}** "
                f"patentamientos ({pct_ajuste:.0f}%) corresponden a Hertz bajo el régimen "
                f"provincial vigente desde 2025, que en su mayoría no circula realmente en "
                f"la provincia; neteando ese efecto, el volumen real queda en "
                f"**{fmt_int(neto_total)}**."
            )

    df2 = df.copy()
    df2["Anio"] = df2["Periodo"].dt.year
    df2["MesNum"] = df2["Periodo"].dt.month
    anios = sorted(df2["Anio"].unique())
    if len(anios) >= 2:
        meses_comunes = sorted(
            set.intersection(*[set(df2[df2["Anio"] == a]["MesNum"]) for a in anios])
        )
        if meses_comunes:
            comp = df2[df2["MesNum"].isin(meses_comunes)]
            por_anio = comp.groupby("Anio")["Valor"].sum()
            a0, a1 = anios[0], anios[-1]
            if por_anio.get(a0, 0) > 0:
                var_ia = (por_anio[a1] - por_anio[a0]) / por_anio[a0] * 100
                direccion = "un incremento" if var_ia > 0.5 else "una caída" if var_ia < -0.5 else "estabilidad"
                partes.append(
                    f"Comparando los mismos meses de {a0} y {a1}, se observa "
                    f"{direccion}{'' if direccion == 'estabilidad' else f' interanual del {abs(var_ia):.0f}%'}."
                )

    return " ".join(partes)


# ---------------------------------------------------------------------------
# Análisis más profundo con IA (Claude) — opcional, requiere API key propia
# ---------------------------------------------------------------------------

def get_anthropic_api_key():
    """Busca la API key de Anthropic en los Secrets de Streamlit. Devuelve
    None si no está configurada (no lanza error)."""
    try:
        return st.secrets.get("ANTHROPIC_API_KEY")
    except Exception:
        return None


def get_gemini_api_key():
    """Busca la API key de Gemini (Google) en los Secrets de Streamlit.
    Devuelve None si no está configurada (no lanza error)."""
    try:
        return st.secrets.get("GEMINI_API_KEY")
    except Exception:
        return None


def _build_stats_context(df: pd.DataFrame, df_bruto: pd.DataFrame = None, df_neto: pd.DataFrame = None) -> str:
    """Arma un bloque de texto compacto con los agregados clave, para
    mandarle a la IA en vez de los datos crudos (más rápido, más barato,
    y no expone el detalle fila por fila)."""
    total = int(df["Valor"].sum())
    periodos = sorted(df["Periodo"].unique())
    rango = f"{month_label(periodos[0])} a {month_label(periodos[-1])}" if periodos else "sin datos"

    monthly = df.groupby("Periodo")["Valor"].sum().sort_index()
    monthly_txt = ", ".join(f"{month_label(p)}: {int(v)}" for p, v in monthly.items())

    def top_txt(col, n=8):
        if total == 0:
            return "sin datos"
        g = df.groupby(col)["Valor"].sum().sort_values(ascending=False).head(n)
        return ", ".join(f"{k}: {int(v)} ({v / total * 100:.1f}%)" for k, v in g.items())

    lines = [
        f"Período analizado: {rango} ({len(periodos)} meses)",
        f"Total patentamientos: {total}",
        f"Evolución mensual: {monthly_txt}",
        f"Por provincia: {top_txt('provincia_label', 10)}",
        f"Por tipo de vehículo: {top_txt('tipo', 10)}",
        f"Top marcas: {top_txt('Marca', 8)}",
        f"Top modelos: {top_txt('Modelo', 8)}",
        f"Top localidades: {top_txt('Localidad', 8)}",
    ]

    if df_bruto is not None and df_neto is not None and not df_bruto.empty:
        bruto_total = int(df_bruto["Valor"].sum())
        neto_total = int(df_neto["Valor"].sum())
        ajuste = bruto_total - neto_total
        lines.append(
            f"IMPORTANTE — contexto de negocio: desde marzo/abril de 2025, la "
            f"empresa Hertz patenta vehículos en la provincia bajo un régimen "
            f"provincial que le resulta beneficioso, pero la mayoría de esos "
            f"autos no circula realmente en la provincia. Volumen global "
            f"(bruto, incluye Hertz): {bruto_total}. Volumen real (neto, "
            f"descontando Hertz): {neto_total}. Diferencia atribuida a Hertz: "
            f"{ajuste} ({ajuste / bruto_total * 100 if bruto_total else 0:.1f}% del "
            f"bruto). Si hay picos de patentamientos a partir de marzo/abril "
            f"2025, considerá que pueden estar influidos por este efecto y no "
            f"necesariamente por demanda genuina del mercado."
        )

    df2 = df.copy()
    df2["Anio"] = df2["Periodo"].dt.year
    df2["MesNum"] = df2["Periodo"].dt.month
    anios = sorted(df2["Anio"].unique())
    if len(anios) >= 2:
        meses_comunes = sorted(
            set.intersection(*[set(df2[df2["Anio"] == a]["MesNum"]) for a in anios])
        )
        if meses_comunes:
            comp = df2[df2["MesNum"].isin(meses_comunes)]
            por_anio = comp.groupby("Anio")["Valor"].sum()
            lines.append(
                "Comparación interanual (mismos meses): "
                + ", ".join(f"{a}: {int(v)}" for a, v in por_anio.items())
            )

    return "\n".join(lines)


def _build_system_prompt() -> str:
    return (
        "Sos un analista de datos del mercado automotor de la Patagonia "
        "argentina (Chubut, Río Negro y Santa Cruz). A partir de los números "
        "de patentamientos que te paso, escribí un análisis breve (180 a 250 "
        "palabras), en español rioplatense, con tono profesional y cercano. "
        "Destacá tendencias, cambios relevantes, concentraciones o "
        "desequilibrios entre provincias, marcas o tipos de vehículo, y "
        "alguna pregunta o hipótesis que valga la pena investigar más. Si el "
        "contexto menciona el efecto Hertz, tenelo en cuenta al interpretar "
        "picos o caídas — no lo confundas con demanda genuina del mercado. "
        "Usá EXCLUSIVAMENTE los datos provistos abajo, no inventes cifras "
        "que no estén ahí. Escribí en prosa corrida, sin títulos ni listas."
    )


def generate_ai_analysis(df: pd.DataFrame, api_key: str, df_bruto: pd.DataFrame = None,
                          df_neto: pd.DataFrame = None, model: str = "claude-sonnet-5") -> str:
    """Llama a la API de Anthropic (Claude) con un resumen de los datos ya
    agregados y devuelve un párrafo de análisis en español. Puede lanzar
    excepciones (auth, red, etc.) — el caller debe manejarlas."""
    import anthropic

    client = anthropic.Anthropic(api_key=api_key)
    context = _build_stats_context(df, df_bruto, df_neto)
    resp = client.messages.create(
        model=model,
        max_tokens=700,
        system=_build_system_prompt(),
        messages=[{"role": "user", "content": context}],
    )
    return "".join(block.text for block in resp.content if hasattr(block, "text"))


def generate_ai_analysis_gemini(df: pd.DataFrame, api_key: str, df_bruto: pd.DataFrame = None,
                                 df_neto: pd.DataFrame = None, model: str = "gemini-3.5-flash") -> str:
    """Llama a la API de Gemini (Google) con el mismo resumen de datos y
    devuelve un párrafo de análisis en español. Puede lanzar excepciones
    (auth, red, etc.) — el caller debe manejarlas."""
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)
    context = _build_stats_context(df, df_bruto, df_neto)
    response = client.models.generate_content(
        model=model,
        contents=context,
        config=types.GenerateContentConfig(
            system_instruction=_build_system_prompt(),
            max_output_tokens=800,
        ),
    )
    return response.text
